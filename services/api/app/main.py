from __future__ import annotations

import csv
import io
import sqlite3
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import HTMLResponse
from pydantic import BaseModel, Field

DB_PATH = Path(__file__).resolve().parent.parent / "data.db"
DB_PATH.parent.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="NahaLabs AI Operator API", version="0.1.0")


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def db() -> sqlite3.Connection:
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    return connection


def init_db() -> None:
    with db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS jobs (
                id TEXT PRIMARY KEY,
                instruction TEXT NOT NULL,
                status TEXT NOT NULL,
                profile TEXT NOT NULL,
                device_serial TEXT,
                claimed_by TEXT,
                result_json TEXT,
                error TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS ingested_files (
                id TEXT PRIMARY KEY,
                filename TEXT NOT NULL,
                rows_json TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            """
        )


init_db()


class CreateJob(BaseModel):
    instruction: str = Field(min_length=1, max_length=4000)
    profile: str = Field(default="flash", pattern="^(flash|pro)$")
    device_serial: str | None = None


class ClaimJob(BaseModel):
    worker_id: str = Field(min_length=1, max_length=128)


class CompleteJob(BaseModel):
    worker_id: str = Field(min_length=1, max_length=128)
    result: dict[str, Any] = Field(default_factory=dict)
    error: str | None = None


@app.get("/", response_class=HTMLResponse)
def dashboard() -> str:
    html = Path(__file__).with_name("dashboard.html").read_text(encoding="utf-8")
    return html


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok", "service": "nahalabs-operator-api"}


@app.post("/v1/files/ingest")
async def ingest_file(file: UploadFile = File(...)) -> dict[str, Any]:
    filename = file.filename or "upload"
    raw = await file.read()
    suffix = Path(filename).suffix.lower()

    if suffix == ".csv":
        text = raw.decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(text)))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            rows = []
        else:
            headers = [str(v or "").strip() for v in values[0]]
            rows = [dict(zip(headers, row, strict=False)) for row in values[1:]]
    else:
        raise HTTPException(status_code=400, detail="Only CSV and XLSX are supported")

    file_id = str(uuid.uuid4())
    import json

    with db() as conn:
        conn.execute(
            "INSERT INTO ingested_files VALUES (?, ?, ?, ?)",
            (file_id, filename, json.dumps(rows, default=str), now()),
        )

    return {"id": file_id, "filename": filename, "rows": len(rows), "preview": rows[:10]}


@app.post("/v1/jobs")
def create_job(payload: CreateJob) -> dict[str, Any]:
    job_id = str(uuid.uuid4())
    timestamp = now()
    with db() as conn:
        conn.execute(
            "INSERT INTO jobs VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                job_id,
                payload.instruction,
                "queued",
                payload.profile,
                payload.device_serial,
                None,
                None,
                None,
                timestamp,
                timestamp,
            ),
        )
    return {"id": job_id, "status": "queued", "created_at": timestamp}


@app.get("/v1/jobs")
def list_jobs(limit: int = 50) -> list[dict[str, Any]]:
    limit = max(1, min(limit, 200))
    with db() as conn:
        rows = conn.execute("SELECT * FROM jobs ORDER BY created_at DESC LIMIT ?", (limit,)).fetchall()
    return [dict(row) for row in rows]


@app.get("/v1/jobs/{job_id}")
def get_job(job_id: str) -> dict[str, Any]:
    with db() as conn:
        row = conn.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Job not found")
    item = dict(row)
    if item["result_json"]:
        import json

        item["result"] = json.loads(item.pop("result_json"))
    else:
        item.pop("result_json")
    return item


@app.post("/v1/jobs/claim")
def claim_job(payload: ClaimJob) -> dict[str, Any] | None:
    with db() as conn:
        row = conn.execute(
            "SELECT * FROM jobs WHERE status = 'queued' ORDER BY created_at LIMIT 1"
        ).fetchone()
        if row is None:
            return None
        conn.execute(
            "UPDATE jobs SET status='running', claimed_by=?, updated_at=? WHERE id=? AND status='queued'",
            (payload.worker_id, now(), row["id"]),
        )
        row = conn.execute("SELECT * FROM jobs WHERE id=?", (row["id"],)).fetchone()
    return dict(row)


@app.post("/v1/jobs/{job_id}/complete")
def complete_job(job_id: str, payload: CompleteJob) -> dict[str, Any]:
    import json

    with db() as conn:
        row = conn.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Job not found")
        if row["claimed_by"] != payload.worker_id:
            raise HTTPException(status_code=409, detail="Job is not owned by this worker")
        status = "failed" if payload.error else "succeeded"
        conn.execute(
            "UPDATE jobs SET status=?, result_json=?, error=?, updated_at=? WHERE id=?",
            (status, json.dumps(payload.result, default=str), payload.error, now(), job_id),
        )
    return {"id": job_id, "status": status}
